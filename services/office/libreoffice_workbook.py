"""Free development renderer; structured OOXML is recalculated by real Calc."""
import json
from pathlib import Path
import subprocess
import sys
from decimal import Decimal
import tempfile
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.packaging.custom import StringProperty
import pymupdf
from workbook_contract import TEMPLATE_VERSION, SHEETS, KEYS, NAMES, digest, validate_input
from calc_engine import document, save_xlsx, export_pdf, export_native_previews


def text(cell, value):
    cell.value = 'Not applicable' if value is None else str(value)
    cell.data_type = 's'


def values(sheet, row, items):
    for col, value in enumerate(items, 1):
        text(sheet.cell(row, col), value)


def build(data, output):
    validate_input(data)
    output.mkdir(parents=True, exist_ok=True)
    version = subprocess.check_output(['libreoffice', '--version'], text=True).strip()
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.calculation.calcMode = 'auto'
    for key in ['revision_id', 'deliverable_id', 'deal_id', 'template_version', 'evaluation_time']:
        workbook.custom_doc_props.append(StringProperty(name=key, value=data[key]))
    workbook.custom_doc_props.append(StringProperty(name='acceptance_profile', value='development_foss_v1'))
    for name in SHEETS:
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = 'B8'
        for col in range(1, 9):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 26 if col == 1 else 18
        for row, value in enumerate([f'DEAL CONTROL / {name.upper()}', data['deal_name'],
            f"Revision {data['revision_id']}", f"{data['purpose']} | {data['audience']}",
            f"{data['confidentiality'].upper()} | {data['provenance'].upper()} | No external-use authorization"], 1):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            text(sheet.cell(row, 1), value)
            sheet.row_dimensions[row].height = 30 if row == 1 else 23
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = '1:7'
        sheet.page_margins.top = sheet.page_margins.bottom = .4
        sheet.page_margins.left = sheet.page_margins.right = .3
        sheet.oddFooter.left.text = f"{data['confidentiality'].upper()} / {data['provenance']}"
        sheet.oddFooter.center.text = f"Rev {data['revision_id']}"
        sheet.oddFooter.right.text = 'Page &P of &N'
    headings = {
        'Overview':['Control','Exact value'],
        'Inputs':['Definition','Value','Period','Unit / currency','Sign / precision','Actual / forecast','Source locator','Authority'],
        'Assumptions':['Definition','Value','Period','Unit / currency','Sign / precision','Actual / forecast','Source locator','Authority'],
        'Valuation':['Scenario','Enterprise value','Cash','Debt','Equity value','Expected equity','Tie-out difference','Units / period'],
        'Scenarios':['Scenario','Equity value','Run / exact model','Assumption scope'],
        'Lineage':['Native region','Source record','Representation','Source SHA-256','Source locator','Fact / assumption','Decision','Run / model / scenario'],
        'Banker Notes':['Ownership','Notes']}
    for name, labels in headings.items():
        values(workbook[name], 7, labels)
        workbook[name].row_dimensions[7].height = 28
    controls = [('Purpose',data['purpose']),('Audience',data['audience']),('Calculation engine',version),
        ('Template',TEMPLATE_VERSION),('Calculation mode','Automatic; recalculated on exact delivered bytes'),
        ('Evaluation time',data['evaluation_time']),('Scope','Controlled EV-to-equity models and scenarios only'),
        ('Readiness','Inspect current Review/QC; generation never establishes readiness'),
        ('Office compatibility','Development LibreOffice profile; Microsoft Excel not verified'),
        ('External use','Not authorized'),('Limitations',' | '.join(data['limitations']))]
    for row, item in enumerate(controls, 8):
        workbook['Overview'].merge_cells(start_row=row,start_column=2,end_row=row,end_column=8)
        values(workbook['Overview'],row,item)
        workbook['Overview'].row_dimensions[row].height=30
    notes=workbook['Banker Notes']
    notes.merge_cells('B8:H11')
    values(notes,8,['Banker-owned','Banker-owned notes; retained on supported save/reopen paths.'])
    notes['B8'].comment=Comment('Banker content must never be silently overwritten.','Deal Control')
    counters={'Inputs':8,'Assumptions':8,'Lineage':8}
    lineage=[]
    for index, run in enumerate(data['calculations'],1):
        row=index+7
        for key,name in zip(KEYS,NAMES):
            measure=next(m for m in run['measures'] if m['key']==key)
            sheet_name='Assumptions' if measure['assumption_id'] else 'Inputs'
            sheet=workbook[sheet_name]; source_row=counters[sheet_name];counters[sheet_name]+=1
            locator=json.dumps(measure['locator'],sort_keys=True)
            authority=measure['fact_id'] or measure['assumption_id']
            values(sheet,source_row,[measure['definition'],'',measure['period'],f"{measure['unit']} / {measure['currency']}",
                f"{measure['sign']} / {measure['precision']}",measure['actual_forecast'],locator,authority])
            cell=sheet.cell(source_row,2);cell.value=float(measure['value'])
            cell.number_format='0'+('.'+'0'*measure['precision'] if measure['precision'] else '')
            cell.comment=Comment(f"Source {measure['source_record_id']} | {locator} | Decision {measure['decision_id']}",'Deal Control')
            sheet.row_dimensions[source_row].height=52
            workbook.defined_names.add(DefinedName(f'{name}_{index}',attr_text=f"'{sheet_name}'!$B${source_row}"))
            lineage.append({**measure,'native_sheet':sheet_name,'native_range':f'B{source_row}','output_sheet':'Valuation',
                'output_range':f'B{row}:G{row}','revision_id':data['revision_id'],'run_id':run['run_id'],
                'calculation_version_id':run['calculation_version_id'],'model_version_id':run['model_version_id'],'scenario_version_id':run['scenario_version_id']})
            line_row=counters['Lineage'];counters['Lineage']+=1
            values(workbook['Lineage'],line_row,[f'{sheet_name}!B{source_row} -> Valuation!B{row}:G{row}',measure['source_record_id'],
                measure['representation_id'],measure['source_digest'],locator,authority,measure['decision_id'],
                f"{run['run_id']} / {run['model_version_id']} / {run['scenario_version_id']}"])
            workbook['Lineage'].row_dimensions[line_row].height=160
        sheet=workbook['Valuation'];text(sheet.cell(row,1),run['scenario'])
        precision=max(m['precision'] for m in run['measures'])
        for col, formula in enumerate([f'=EV_{index}',f'=Cash_{index}',f'=Debt_{index}',f'=ROUND(EV_{index}+Cash_{index}-Debt_{index},{precision})'],2):
            sheet.cell(row,col).value=formula
        sheet.cell(row,6).value=float(run['expected_equity_value'])
        sheet.cell(row,7).value=f'=ROUND(E{row}-F{row},{precision})'
        text(sheet.cell(row,8),f"{run['measures'][0]['unit']} / {run['measures'][0]['period']}")
        for col in range(2,8):sheet.cell(row,col).number_format='0.0#########'
        sheet.row_dimensions[row].height=44
        values(workbook['Scenarios'],row,[run['scenario'],'',f"{run['run_id']} / {run['model_version_id']}",
            'Assumptions remain assumptions; scenario is not a prediction.'])
        workbook['Scenarios'].cell(row,2).value=f'=Valuation!E{row}'
        workbook['Scenarios'].row_dimensions[row].height=44
    chart=BarChart();chart.title='Equity value by controlled scenario';chart.type='col'
    chart.add_data(Reference(workbook['Scenarios'],min_col=2,min_row=7,max_row=7+len(data['calculations'])),titles_from_data=True)
    chart.set_categories(Reference(workbook['Scenarios'],min_col=1,min_row=8,max_row=7+len(data['calculations'])))
    chart.series[0].graphicalProperties.solidFill='4472C4'
    chart.plot_area.spPr=GraphicalProperties(solidFill='C0C0C0')
    chart.y_axis.title=data['calculations'][0]['measures'][0]['unit']
    chart.x_axis.numFmt=NumFmt(formatCode='General',sourceLinked=False)
    chart.y_axis.numFmt=NumFmt(formatCode='0.0#########',sourceLinked=False)
    chart.width,chart.height=25,10
    workbook['Scenarios'].add_chart(chart,f'A{len(data["calculations"])+10}')
    for sheet in workbook:
        last=max(18,sheet.max_row)
        if sheet.title=='Scenarios':last=len(data['calculations'])+30
        sheet.print_area=f'A1:H{last}'
        for row in sheet.iter_rows(max_row=last,max_col=8):
            for cell in row:
                cell.font=Font(name='DejaVu Sans',size=10,color='18352C')
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.row==7:
                    cell.fill=PatternFill('solid',fgColor='E8EEEA')
                    cell.font=Font(name='DejaVu Sans',size=10,color='18352C',bold=True)
        sheet['A1'].fill=PatternFill('solid',fgColor='18352C')
        sheet['A1'].font=Font(name='DejaVu Sans',size=14,color='FFFFFF',bold=True)
    with tempfile.TemporaryDirectory(prefix='calc-native-') as temporary:
        initial=Path(temporary)/'input.xlsx';workbook.save(initial)
        native=output/'analysis-valuation.xlsx'
        with document(initial) as actual:
            # Ask the actual font/layout engine to size unmerged data rows before
            # freezing the delivered bytes; fixed heights can clip long citations.
            ends={'Inputs':counters['Inputs']-1,'Assumptions':counters['Assumptions']-1,
                'Lineage':counters['Lineage']-1,'Valuation':7+len(data['calculations']),
                'Scenarios':7+len(data['calculations'])}
            for name,last_row in ends.items():
                if last_row>=8:
                    actual.Sheets.getByName(name).getCellRangeByName(f'A8:H{last_row}').Rows.OptimalHeight=True
            save_xlsx(actual,native)
        # The PDF is rendered only after reopening the exact delivered XLSX.
        with document(native) as exact:export_pdf(exact,output/'analysis-valuation.pdf')
    cached=openpyxl.load_workbook(native,data_only=True)
    formulas=openpyxl.load_workbook(native,data_only=False)
    checks=[]
    for row,run in enumerate(data['calculations'],8):
        expected=Decimal(run['expected_equity_value']);value=Decimal(str(cached['Valuation'][f'E{row}'].value))
        delta=Decimal(str(cached['Valuation'][f'G{row}'].value))
        checks.append({'locator':f'Valuation!E{row}:G{row}','expected':str(expected),'observed':str(value),'difference':str(delta),'passed':value==expected and delta==0})
    pages=[]
    native_pages=export_native_previews(native,output,[(s.title,f'A1:H{s.max_row}') for s in formulas])
    with pymupdf.open(output/'analysis-valuation.pdf') as pdf:
        for index,page in enumerate(pdf,1):
            content=page.get_text();sheet=next((name for name in SHEETS if f'DEAL CONTROL / {name.upper()}' in content),None)
            image=f'reader-page-{index}.png';page.get_pixmap(matrix=pymupdf.Matrix(1.2,1.2)).save(output/image)
            pages.append({'page':index,'image':image,'text':content,'width':page.rect.width,'height':page.rect.height})
    for relation in lineage:
        relation['reader_pages']=[p['page'] for p in pages if any(f'DEAL CONTROL / {name.upper()}' in p['text'] for name in [relation['native_sheet'],'Valuation','Scenarios','Lineage'])]
    inventory=[{'sheet':s.title,'cells':[{'cell':c.coordinate,'formula':c.value if c.data_type=='f' else None,'value':cached[s.title][c.coordinate].value,
        'number_format':c.number_format,'comment':c.comment.text if c.comment else None} for row in s for c in row if c.value is not None]} for s in formulas]
    report={'schema_version':'1.0.0','revision_id':data['revision_id'],'engine':'libreoffice.calc','engine_version':version,'template_version':TEMPLATE_VERSION,
        'acceptance_profile':'development_foss_v1','licensed':True,'font':'DejaVu Sans',
        'font_manifest':[{'file':p.name,'sha256':digest(p.read_bytes())} for p in sorted(Path('/usr/share/fonts/truetype/dejavu').glob('*.ttf'))],
        'calculation_mode':'automatic','evaluation_time':data['evaluation_time'],'recalculation':{'outcome':'passed' if all(c['passed'] for c in checks) else 'failed','checks':checks},
        'files':[{'path':name,'sha256':digest((output/name).read_bytes()),'bytes':(output/name).stat().st_size} for name in ['analysis-valuation.xlsx','analysis-valuation.pdf']],
        'pages':pages,'native_pages':native_pages,'cells':inventory,'lineage':lineage,
        'limitations':['Development LibreOffice profile; Microsoft Excel compatibility is not established.','Native worksheet images are exported directly from the exact XLSX selection; Reader images are rasterized independently from the PDF.']}
    (output/'render-report.json').write_text(json.dumps(report,indent=2))
    return report

if __name__=='__main__':
    import os
    os.environ['OFFICE_ENGINE']='libreoffice'
    data=json.loads(Path(sys.argv[1]).read_text());output=Path(sys.argv[2]);build(data,output)
    from inspect_workbook import inspect
    (output/'inspection-report.json').write_text(json.dumps(inspect(data,output/'analysis-valuation.xlsx',output/'analysis-valuation.pdf'),indent=2))
