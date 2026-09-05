"""Independent file observers: openpyxl + PDF parser; never regenerates the input."""
import hashlib
import json
import re
from decimal import Decimal
from pathlib import Path
import sys
import zipfile
import openpyxl
import aspose.cells as cells
import pymupdf

SHEETS = ['Overview', 'Inputs', 'Assumptions', 'Valuation', 'Scenarios', 'Lineage', 'Banker Notes']
KEYS = [('enterprise_value', 'EV'), ('cash', 'Cash'), ('debt', 'Debt')]

def number(text):
    return Decimal(text.replace(',', '')) if re.fullmatch(r'-?\d[\d,]*(?:\.\d+)?', text) else None

def valuation_rows_match(pages, runs):
    rows=[]
    for page,_ in pages:
        words=page.get_text('words')
        heading=next((w for w in words if w[4]=='Scenario'),None)
        if not heading:return False
        columns=sorted([d['rect'] for d in page.get_drawings() if d.get('fill') and d['rect'].y0<=heading[1]<d['rect'].y1 and d['rect'].width>30],key=lambda r:r.x0)
        if len(columns)!=8:return False
        anchors=sorted([w for w in words if columns[1].x0<=w[0]<columns[1].x1 and w[1]>columns[1].y1 and number(w[4]) is not None],key=lambda w:w[1])
        for index,anchor in enumerate(anchors):
            bottom=anchors[index+1][1]-2 if index+1<len(anchors) else page.rect.height-40
            values=[]
            for column in columns[1:7]:
                tokens=[number(w[4]) for w in words if column.x0<=w[0]<column.x1 and abs(w[1]-anchor[1])<3 and number(w[4]) is not None]
                if len(tokens)!=1:return False
                values.append(tokens[0])
            texts=[' '.join(w[4] for w in words if column.x0<=w[0]<column.x1 and anchor[1]-2<=w[1]<bottom) for column in [columns[0],columns[7]]]
            rows.append((values,texts))
    if len(rows)!=len(runs):return False
    for (values,texts),run in zip(rows,runs):
        measures={m['key']:m for m in run['measures']}
        expected=[Decimal(measures[key]['value']) for key,_ in KEYS]+[Decimal(run['expected_equity_value'])]*2+[Decimal(0)]
        if values!=expected or ' '.join(run['scenario'].split()) not in texts[0]:return False
        if not all(measures['cash'][key] in texts[1] for key in ['unit','period']):return False
    return True

def chart_matches(pages,runs):
    for page,content in pages:
        titles=page.search_for('Equity value by controlled scenario')
        if not titles:continue
        drawings=page.get_drawings()
        plots=[d['rect'] for d in drawings if d.get('fill') and all(abs(c-.7529)<.01 for c in d['fill']) and d['rect'].y0>titles[0].y0 and d['rect'].width>100]
        if len(plots)!=1:continue
        plot=plots[0]
        bars=sorted([d['rect'] for d in drawings if d.get('fill') and all(abs(c-e)<.01 for c,e in zip(d['fill'],(.2667,.4471,.7686))) and plot.contains(d['rect'])],key=lambda r:r.x0)
        words=page.get_text('words')
        ticks=[(number(w[4]),(w[1]+w[3])/2) for w in words if w[2]<plot.x0 and plot.y0-10<w[1]<plot.y1+10 and number(w[4]) is not None]
        if len(ticks)<2:continue
        low,high=min(ticks),max(ticks)
        if low[0]==high[0] or low[1]==high[1]:continue
        scale=float(high[0]-low[0])/(high[1]-low[1])
        zero_y=low[1]-float(low[0])/scale
        matched=0
        for index,run in enumerate(runs):
            expected=Decimal(run['expected_equity_value'])
            left=plot.x0+index*plot.width/len(runs);right=left+plot.width/len(runs)
            candidates=[bar for bar in bars if left<(bar.x0+bar.x1)/2<right]
            if expected==0:
                if any(bar.height>1 for bar in candidates):break
            else:
                if len(candidates)!=1:break
                bar=candidates[0];matched+=1
                endpoint=bar.y0 if expected>0 else bar.y1
                observed=float(low[0])+(endpoint-low[1])*scale
                tolerance=max(.05,abs(scale)*1.5)
                if abs(observed-float(expected))>tolerance:break
            category=' '.join(w[4] for w in words if left<((w[0]+w[2])/2)<right and zero_y<w[1]<zero_y+50)
            if ' '.join(run['scenario'].split()) not in category:break
        else:
            return matched==len([bar for bar in bars if bar.height>1]) and runs[0]['measures'][0]['unit'] in content
    return False


def inspect(data, native, reader):
    checks = []
    observations = {}
    def record(code, passed, detail, locator=None):
        checks.append({'code': code, 'outcome': 'passed' if passed else 'failed', 'detail': detail, 'locator': locator or {}})
    try:
        if native.stat().st_size > 100*1024*1024:
            raise ValueError('native_byte_limit')
        with zipfile.ZipFile(native) as archive:
            if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > 160*1024*1024:
                raise ValueError('native_expansion_limit')
            unsupported = [n for n in archive.namelist() if any(x in n for x in ['vbaProject', 'externalLinks/', 'embeddings/', 'connections.xml'])]
            if unsupported:
                raise ValueError('unsupported_active_content')
        workbook = openpyxl.load_workbook(native, data_only=False, keep_links=False)
        cached = openpyxl.load_workbook(native, data_only=True, keep_links=False)
        props = {p.name: p.value for p in workbook.custom_doc_props}
        structure = all(name in workbook.sheetnames for name in SHEETS) and props.get('revision_id') == data['revision_id']
        structure &= len(workbook['Scenarios']._charts) == 1 and workbook.calculation.calcMode in (None, 'auto')
        series=workbook['Scenarios']._charts[0].series
        reference=lambda value: value.replace('$','').replace("'",'')
        structure &= len(series)==1 and reference(series[0].val.numRef.f)==f'Scenarios!B8:B{7+len(data["calculations"])}'
        category=series[0].cat.strRef or series[0].cat.numRef
        structure &= reference(category.f)==f'Scenarios!A8:A{7+len(data["calculations"])}'
        cache_pass, lineage_pass = True, True
        for index, run in enumerate(data['calculations'], 1):
            row = index+7
            precision=max(m['precision'] for m in run['measures'])
            expected_formula = f'=ROUND(EV_{index}+Cash_{index}-Debt_{index},{precision})'
            structure &= workbook['Valuation'][f'E{row}'].value == expected_formula
            structure &= all(workbook['Valuation'].cell(row,col).value==f'={name}_{index}' for col,(_,name) in enumerate(KEYS,2))
            structure &= workbook['Scenarios'][f'A{row}'].value==run['scenario'] and workbook['Scenarios'][f'B{row}'].value==f'=Valuation!E{row}'
            expected = Decimal(next(m['value'] for m in run['measures'] if m['key']=='enterprise_value')) + Decimal(next(m['value'] for m in run['measures'] if m['key']=='cash')) - Decimal(next(m['value'] for m in run['measures'] if m['key']=='debt'))
            cache_pass &= Decimal(str(cached['Valuation'][f'E{row}'].value)) == expected and Decimal(str(cached['Valuation'][f'G{row}'].value)) == 0
            cache_pass &= Decimal(str(cached['Scenarios'][f'B{row}'].value)) == expected
            for key, name in KEYS:
                measure = next(m for m in run['measures'] if m['key']==key)
                defined = workbook.defined_names.get(f'{name}_{index}')
                if defined is None:
                    structure = False
                    continue
                destinations = list(defined.destinations)
                if len(destinations) != 1 or destinations[0][0] not in ['Inputs','Assumptions']:
                    structure = False
                    continue
                sheet, cell = destinations[0]
                input_cell = workbook[sheet][cell]
                locator = json.dumps(measure['locator'], sort_keys=True)
                lineage_pass &= input_cell.comment is not None and locator in input_cell.comment.text and measure['decision_id'] in input_cell.comment.text
                lineage_pass &= workbook[sheet].cell(input_cell.row, 7).value == locator
                lineage_pass &= Decimal(str(input_cell.value)) == Decimal(measure['value'])
                structure &= sheet == ('Assumptions' if measure['assumption_id'] else 'Inputs')
        license_path=Path('/run/secrets/aspose-license')
        if license_path.is_file():
            cells.License().set_license(str(license_path))
        recalculated=cells.Workbook(str(native))
        recalculated.calculate_formula()
        for index,run in enumerate(data['calculations'],8):
            cache_pass &= Decimal(str(recalculated.worksheets.get('Valuation').cells.get(f'E{index}').double_value)) == Decimal(str(cached['Valuation'][f'E{index}'].value))
        record('native_structure', bool(structure), 'Native formula, defined-name, chart, revision and calculation-mode checks')
        record('recalculation', bool(cache_pass), 'Exact stored Native reopened and recalculated by Aspose; independent decimals compared with stored caches and chart data')
        record('lineage', bool(lineage_pass), 'Input authority, source locator and exact native cell checked independently')
    except Exception as error:
        record('native_structure', False, f'Native Artifact inspection failed: {type(error).__name__}')
        record('recalculation', False, 'Cannot verify stored caches on an invalid Native Artifact')
        record('lineage', False, 'Cannot verify lineage on an invalid Native Artifact')
    try:
        pdf = pymupdf.open(reader)
        text = '\n'.join(p.get_text() for p in pdf)
        clean = not re.search(r'evaluation\s+only|evaluation\s+warning|created\s+(?:with|by)\s+aspose|aspose[^\n]{0,80}(?:trial|unlicensed)', text, re.I)
        record('clean_copy', clean, 'Reader Copy checked for evaluation marks and clean output')
        pages = [(page, page.get_text()) for page in pdf]
        sheet_pages = {sheet: [(page, content) for page, content in pages if f'DEAL CONTROL / {sheet.upper()}' in content] for sheet in SHEETS}
        order = []
        for _, content in pages:
            for sheet in SHEETS:
                if f'DEAL CONTROL / {sheet.upper()}' in content and (not order or order[-1] != sheet):
                    order.append(sheet)
        observations['sheet_order_matches'] = order == SHEETS
        observations['valuation_values_match'] = valuation_rows_match(sheet_pages['Valuation'],data['calculations'])
        normalized_lineage = re.sub(r'\s+', '', '\n'.join(content for _, content in sheet_pages['Lineage']))
        citations_match = all(m['decision_id'] in normalized_lineage and (not m.get('source_record_id') or m['source_record_id'] in normalized_lineage) for run in data['calculations'] for m in run['measures'])
        observations['decision_and_source_citations_match'] = citations_match
        chart_pages = sheet_pages['Scenarios']
        observations['chart_region_present'] = chart_matches(chart_pages,data['calculations'])
        parity = all(observations.values()) and data['revision_id'] in text and data['confidentiality'].upper() in text
        parity &= data['purpose'] in text and data['audience'] in text
        fonts = [f for page in pdf for f in page.get_fonts(full=True)]
        embedded = bool(fonts) and all(bool(pdf.extract_font(font[0])[3]) for font in fonts)
        safe = not pdf.is_form_pdf and pdf.embfile_count() == 0 and all('JavaScript' not in pdf.xref_object(i) and '/Launch' not in pdf.xref_object(i) for i in range(1,pdf.xref_length()))
        observations['fonts_embedded'] = embedded
        observations['safe_pdf_17'] = safe and pdf.metadata.get('format') == 'PDF 1.7'
        record('native_reader_parity', bool(parity and all(observations.values())), 'Sheet order/content, revision, numbers, confidentiality, embedded fonts and PDF active content checked; visual review remains explicit')
    except Exception as error:
        record('native_reader_parity', False, f'Reader Copy inspection failed: {type(error).__name__}')
    unknown = any(m['actual_forecast']=='unknown' for r in data['calculations'] for m in r['measures'])
    checks.append({'code':'controlled_inputs','outcome':'missing' if unknown else 'passed','detail':'Actual/forecast classification incomplete' if unknown else 'Exact controlled inputs and explicit assumptions are pinned'})
    return {'schema_version':'1.0.0','revision_id':data['revision_id'],'native_sha256':hashlib.sha256(native.read_bytes()).hexdigest(),
            'reader_sha256':hashlib.sha256(reader.read_bytes()).hexdigest(),'checks':checks,'observations':observations}


if __name__ == '__main__':
    result = inspect(json.loads(Path(sys.argv[1]).read_text()), Path(sys.argv[2]), Path(sys.argv[3]))
    Path(sys.argv[4]).write_text(json.dumps(result, indent=2))
