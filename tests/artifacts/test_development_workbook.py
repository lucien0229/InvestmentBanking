"""Independent exact-file acceptance for the explicit free development profile."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import openpyxl
import hashlib
import pymupdf
import re

ROOT = Path(__file__).resolve().parents[2]

class DevelopmentWorkbook(unittest.TestCase):
    def test_clean_native_reader_and_recalculation(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            result = subprocess.run([sys.executable, str(ROOT/'services/office/libreoffice_workbook.py'),
                str(ROOT/'tests/fixtures/analysis-workbook.json'), str(root)], capture_output=True, text=True, timeout=150)
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads((root/'inspection-report.json').read_text())
            self.assertTrue(all(check['outcome']=='passed' for check in report['checks']), report)
            formula = openpyxl.load_workbook(root/'analysis-valuation.xlsx', data_only=False)
            cached = openpyxl.load_workbook(root/'analysis-valuation.xlsx', data_only=True)
            self.assertEqual(formula['Valuation']['E8'].value, '=ROUND(EV_1+Cash_1-Debt_1,1)')
            self.assertEqual(cached['Valuation']['E8'].value, 94.7)
            self.assertEqual(cached['Valuation']['G8'].value, 0)
            self.assertEqual(len(formula['Scenarios']._charts), 1)
            self.assertIsNotNone(formula['Banker Notes']['B8'].comment)
            render = json.loads((root/'render-report.json').read_text())
            self.assertEqual(render['engine'], 'libreoffice.calc')
            self.assertEqual(render['acceptance_profile'], 'development_foss_v1')
            with pymupdf.open(root/'analysis-valuation.pdf') as pdf:
                self.assertTrue(all(not list(page.annots() or []) for page in pdf), 'Reader annotations must not cover periods or values')
                scenario_page=next(page for page in pdf if 'DEAL CONTROL / SCENARIOS' in page.get_text())
                visible=re.sub(r'\s+', '', scenario_page.get_text())
                fixture=json.loads((ROOT/'tests/fixtures/analysis-workbook.json').read_text())
                for run in fixture['calculations']:
                    self.assertIn(run['model_version_id'], visible, 'Full model citation must be visible, not clipped by row height')
            previews=render['native_pages']
            self.assertEqual(len(previews),7)
            native_hashes={hashlib.sha256((root/p['image']).read_bytes()).hexdigest() for p in previews}
            reader_hashes={hashlib.sha256((root/p['image']).read_bytes()).hexdigest() for p in render['pages']}
            self.assertEqual(len(native_hashes),7)
            self.assertFalse(native_hashes & reader_hashes)
            self.assertTrue(all(p['render_basis']=='exact_xlsx_cell_selection' for p in previews))
            lab=subprocess.run([sys.executable,str(ROOT/'services/office/compatibility_lab.py'),str(ROOT/'tests/fixtures/analysis-workbook.json'),
                str(root/'analysis-valuation.xlsx'),str(root/'analysis-valuation.pdf'),str(root/'roundtrip')],capture_output=True,text=True,timeout=90)
            self.assertEqual(lab.returncode,0,lab.stderr + ((root/'roundtrip/office-lab-report.json').read_text() if (root/'roundtrip/office-lab-report.json').exists() else ''))
            receipt=json.loads((root/'roundtrip/office-lab-report.json').read_text())
            self.assertEqual(receipt['outcome'],'passed',receipt)
            self.assertNotEqual(receipt['native_sha256'],receipt['returned_sha256'])

if __name__ == '__main__':
    unittest.main()
