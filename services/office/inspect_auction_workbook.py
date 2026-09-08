import hashlib, json, zipfile
from pathlib import Path
import openpyxl, fitz

def inspect(data,native,reader):
    checks=[]
    def add(code,ok,detail): checks.append({"code":code,"outcome":"passed" if ok else "failed","detail":detail,"locator":{}})
    try:
        wb=openpyxl.load_workbook(native,data_only=False,keep_links=False)
        cached=openpyxl.load_workbook(native,data_only=True,keep_links=False)
        props={p.name:p.value for p in wb.custom_doc_props}
        expected_ids=[b["id"] for b in (data.get("process_state") or {}).get("buyers",[])]
        observed_ids=[wb["Buyer Universe"].cell(r,1).value for r in range(9,9+len(expected_ids))]
        add("controlled_inputs",props.get("revision_id")==data.get("revision_id") and props.get("build_input_sha256")==hashlib.sha256(json.dumps(data,sort_keys=True).encode()).hexdigest(),"Exact process snapshot digest and Revision identity are pinned; stale projections fail.")
        add("native_structure",wb.sheetnames[:5]==["Executive Control","Buyer Universe","Process State","Lineage","Banker Notes"] and wb["Executive Control"]["A17"].value=="No aggregate ready / OK score is calculated.","Executive Control is first visible and avoids scalar readiness.")
        add("stable_process_identity",observed_ids==expected_ids and all(wb["Buyer Universe"].cell(r,8).value==identity for r,identity in enumerate(expected_ids,9)),"Buyer identities must match authoritative source objects exactly.")
        add("visible_process_rows",all(ws.sheet_state=="visible" and not any(row.hidden for row in ws.row_dimensions.values()) and not any(col.hidden for col in ws.column_dimensions.values()) for ws in wb),"Hidden sheets, rows or columns cannot conceal process state.")
        add("recalculation",wb["Executive Control"]["B18"].value=="=COUNTA('Buyer Universe'!A9:A1000)" and cached["Executive Control"]["B18"].value==len(expected_ids) and "BuyerCandidateIDs" in wb.defined_names,"Exact native formula and stored recalculation agree with source object count.")
        add("banker_protection",bool(wb["Banker Notes"]["B9"].value) and wb["Banker Notes"]["B9"].comment is not None,"Banker Notes retains content and an explicit protected-content comment.")
        pdf=fitz.open(reader); text="\n".join(p.get_text() for p in pdf)
        add("native_reader_parity",data.get("revision_id") in text and "DEAL CONTROL / EXECUTIVE CONTROL" in text and all(identity in text for identity in expected_ids),"Reader contains the same Revision, executive control and stable Buyer identities.")
        add("lineage",all(wb["Lineage"].cell(r,1).value==identity and wb["Lineage"].cell(r,4).value==f"A{r}:H{r}" for r,identity in enumerate(expected_ids,9)),"Stable object identity and exact native range lineage is present.")
        add("process_state_separation",wb["Process State"].max_row>=15,"Current, candidate, approved and not-applicable states have separate rows.")
        with zipfile.ZipFile(native) as z: add("unsupported_active_content",not any("vbaProject" in n or "externalLinks" in n for n in z.namelist()),"No active content or external links.")
    except Exception as exc:
        add("native_structure",False,str(exc))
    return {"revision_id":data.get("revision_id"),"checks":checks,"report":{"revision_id":data.get("revision_id"),"template_version":"auction-control-1.0.0"}}

if __name__ == "__main__":
    import sys
    result=inspect(json.loads(Path(sys.argv[1]).read_text()),Path(sys.argv[2]),Path(sys.argv[3])); Path(sys.argv[4]).write_text(json.dumps(result,indent=2))
