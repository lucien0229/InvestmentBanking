import copy
import json
import tempfile
import unittest
from pathlib import Path

from services.office.inspect_cim_presentation import inspect
from services.office.cim_presentation import build


def fixture():
    evidence = "22222222-2222-4222-8222-222222222222"
    source = "33333333-3333-4333-8333-333333333333"
    fact = "44444444-4444-4444-8444-444444444444"
    assumption = "55555555-5555-4555-8555-555555555555"
    return {
        "schema_version": "1.0.0", "task_definition": "cim_content_draft", "status": "proposal_only",
        "revision_id": "11111111-1111-4111-8111-111111111111", "template_version": "cim-1.0.0",
        "purpose": "Preparation marketing", "audience": "Internal Banker", "confidentiality": "confidential",
        "approved_disclosure_set": ["DISC-001"], "output_ceiling": {"max_slides": 3},
        "evidence": [{"id": evidence, "source_record_id": source}],
        "facts_assumptions": [{"id": fact, "kind": "fact", "status": "accepted"}, {"id": assumption, "kind": "assumption", "status": "approved"}],
        "content_draft": {"task_definition": "cim_content_draft", "status": "proposal_only", "approved_disclosure_set": ["DISC-001"], "output_ceiling": {"max_slides": 3}, "evidence": [{"id": evidence, "source_record_id": source}], "facts_assumptions": [{"id": fact, "kind": "fact", "status": "accepted"}, {"id": assumption, "kind": "assumption", "status": "approved"}], "sections": [{"section_key": "overview", "title": "Overview", "body": "Controlled proposal", "qualification": "Proposal-only; Banker Review required", "citations": ["DISC-001"], "evidence_refs": [evidence], "fact_refs": [fact], "assumption_refs": [assumption], "source_refs": [source], "table_rows": [["Metric", "Value"], ["Revenue", "$94.6m"]]}, {"section_key": "growth", "title": "Growth", "body": "Subject to review", "qualification": "Proposal-only; Banker Review required", "citations": ["DISC-001"], "evidence_refs": [evidence], "fact_refs": [fact], "assumption_refs": [assumption], "source_refs": [source], "chart": {"categories": ["2023A", "2024A"], "values": [40, 60], "series_name": "Revenue"}}]},
    }


class CIMObserverTest(unittest.TestCase):
    def test_render_and_adversarial_native_reader_checks(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data = fixture()
            build(data, root)
            passed = inspect(data, root / "cim.pptx", root / "cim.pdf", json.loads((root / "render-report.json").read_text()))
            self.assertTrue(all(row["outcome"] == "passed" for row in passed["checks"]))
            self.assertEqual(next(row for row in passed["checks"] if row["code"] == "render_identity_qc_scope")["outcome"], "passed")

            import fitz
            reordered = fitz.open()
            original = fitz.open(root / "cim.pdf")
            reordered.insert_pdf(original, from_page=2, to_page=2)
            reordered.insert_pdf(original, from_page=0, to_page=1)
            reordered.save(root / "reordered.pdf")
            result = inspect(data, root / "cim.pptx", root / "reordered.pdf")
            self.assertEqual(next(row for row in result["checks"] if row["code"] == "native_reader_parity")["outcome"], "failed")

            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            presentation = Presentation(str(root / "cim.pptx"))
            for shape in list(presentation.slides[2].shapes):
                if shape.shape_type == MSO_SHAPE_TYPE.CHART:
                    element = shape._element
                    element.getparent().remove(element)
                    break
            presentation.save(root / "missing-chart.pptx")
            result = inspect(data, root / "missing-chart.pptx", root / "cim.pdf")
            self.assertEqual(next(row for row in result["checks"] if row["code"] == "native_editable")["outcome"], "failed")

            hidden = Presentation(str(root / "cim.pptx"))
            hidden.slides[1].notes_slide.notes_text_frame.text = "INTERNAL BANKER NOTE: remove before circulation"
            hidden.save(root / "hidden-note.pptx")
            result = inspect(data, root / "hidden-note.pptx", root / "cim.pdf")
            self.assertEqual(next(row for row in result["checks"] if row["code"] == "hidden_internal_notes")["outcome"], "failed")

    def test_cover_counts_against_output_ceiling(self):
        data = fixture()
        data["content_draft"]["output_ceiling"] = {"max_slides": 2}
        with self.assertRaises(ValueError):
            build(data, Path(tempfile.mkdtemp()))


if __name__ == "__main__":
    unittest.main()
